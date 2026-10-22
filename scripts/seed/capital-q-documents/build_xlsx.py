"""Build Capital Q's demonstration financial model and cap table (XLSX).

Demonstration data only. Every figure that also appears in
supabase/seed-assets/demo-companies/companies.json is reconciled on the
Annual Summary / Summary sheets so a drift shows up as "CHECK".
Run: python3 scripts/seed/capital-q-documents/build_xlsx.py <out_dir>
"""

import sys
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

WATERMARK = (
    "DEMONSTRATION DOCUMENT — illustrative, prepared for the Capital Q product "
    "demo. Not a legal or financial record."
)
FONT = "Arial"
INK, GOLD = "16140F", "C9A227"
BLUE, GREEN = "0000FF", "008000"
F_IN = Font(name=FONT, color=BLUE, size=10)
F_LINK = Font(name=FONT, color=GREEN, size=10)
F_N = Font(name=FONT, size=10)
F_B = Font(name=FONT, size=10, bold=True)
F_H = Font(name=FONT, size=14, bold=True, color=INK)
F_WM = Font(name=FONT, size=10, bold=True, color="FFFFFF")
FILL_WM = PatternFill("solid", fgColor=INK)
FILL_HEAD = PatternFill("solid", fgColor="F1EDE4")
FILL_KEY = PatternFill("solid", fgColor="FFFF00")
THIN = Side(style="thin", color="D9D1C1")
USD = '#,##0;(#,##0);"-"'
USD2 = '#,##0.00;(#,##0.00);"-"'
PCT = '0.0%;(0.0%);"-"'
NUM1 = '#,##0.0;(#,##0.0);"-"'
NUM0 = '#,##0;(#,##0);"-"'


def props(wb, title, subject):
    p = wb.properties
    p.title = f"{title} — DEMONSTRATION DOCUMENT"
    p.subject = subject
    p.creator = "Capital Q product demo (illustrative)"
    p.description = WATERMARK
    p.keywords = "demonstration; illustrative; not a legal or financial record"
    p.category = "Demonstration data"


def sheet_frame(ws, title, width_cols):
    """Row 1 carries the permanent watermark; it is also the print header."""
    ws["A1"] = WATERMARK
    ws["A1"].font = F_WM
    ws["A1"].fill = FILL_WM
    ws["A1"].alignment = Alignment(vertical="center")
    for c in range(2, width_cols + 1):
        ws.cell(row=1, column=c).fill = FILL_WM
    ws.row_dimensions[1].height = 22
    ws["A2"] = title
    ws["A2"].font = F_H
    ws.oddHeader.center.text = "DEMONSTRATION DOCUMENT — illustrative"
    ws.oddFooter.center.text = WATERMARK
    ws.oddFooter.center.size = 7
    ws.freeze_panes = "A4"
    ws.sheet_view.showGridLines = False


def head(ws, row, values, start=1):
    for i, v in enumerate(values):
        c = ws.cell(row=row, column=start + i, value=v)
        c.font = F_B
        c.fill = FILL_HEAD
        c.border = Border(bottom=THIN)
        c.alignment = Alignment(horizontal="center" if i else "left")


# ---------------------------------------------------------------- model
YEARS = [2026, 2027, 2028]
MONTHS = [(2026 + i // 12, i % 12 + 1) for i in range(36)]
MON = "Jan Feb Mar Apr May Jun Jul Aug Sep Oct Nov Dec".split()
ACTUAL_MONTHS = 9  # Jan-Sep 2026 are self-reported actuals
ACT_ACCOUNTS = [0, 0, 0, 3, 5, 8, 9, 10, 12]
ACT_MRR = [0, 0, 0, 1200, 2000, 3000, 3600, 4200, 4800]


def headcount(i):
    """Plan: four engineers hired after the SAFE closes (Apr and Jun 2027)."""
    eng = 2 + 2 * (i >= 15) + 2 * (i >= 17) + (i >= 24) + (i >= 30)
    ops = (1 if i < 3 else 2) + (i >= 18) + (i >= 21) + (i >= 27) + (i >= 33)
    return 2, eng, ops


def build_model(out):
    wb = Workbook()
    props(wb, "Capital Q financial model 2026–2028", "3-year monthly model")

    # ---- Cover
    cv = wb.active
    cv.title = "Cover"
    sheet_frame(cv, "Capital Q — Financial model 2026–2028", 6)
    cv.column_dimensions["A"].width = 100
    lines = [
        "Home of Private Capital · Lagos, Nigeria · pre-seed",
        "",
        "What this is: an illustrative 36-month model (Jan 2026 – Dec 2028) built for the Capital Q product demo.",
        "Jan–Sep 2026 are self-reported demonstration actuals; Oct 2026 onward is forecast. Nothing here is verified.",
        "Figures shared with the Capital Q profile (MRR, ARR, accounts, margins, team, raise) reconcile on 'Annual Summary'.",
        "",
        "Sheets: Assumptions (all inputs) · Monthly Model (formulas) · Annual Summary (roll-up and reconciliation).",
        "Colour key: blue text = hardcoded input · black = formula · green = link to another sheet · yellow fill = key assumption.",
        "Currency: USD throughout. Cash is modelled before tax; no debt.",
    ]
    for i, t in enumerate(lines):
        c = cv.cell(row=4 + i, column=1, value=t)
        c.font = F_N

    # ---- Assumptions
    a = wb.create_sheet("Assumptions")
    sheet_frame(a, "Assumptions (inputs)", 6)
    for col, w in zip("ABCDEF", (52, 14, 14, 14, 14, 60)):
        a.column_dimensions[col].width = w
    head(a, 4, ["Input", "Unit", "2026", "2027", "2028", "Note / source"])
    rows = {}

    def arow(r, key, label, unit, vals, fmt, note, key_assumption=False):
        rows[key] = r
        a.cell(row=r, column=1, value=label).font = F_N
        a.cell(row=r, column=2, value=unit).font = F_N
        for j, v in enumerate(vals):
            if v is None:
                continue
            c = a.cell(row=r, column=3 + j, value=v)
            c.font = F_IN
            c.number_format = fmt
            if key_assumption:
                c.fill = FILL_KEY
        a.cell(row=r, column=6, value=note).font = F_N

    def section(r, t):
        a.cell(row=r, column=1, value=t).font = F_B

    section(5, "Revenue")
    arow(6, "new_h1", "New paying investor firms per month, Jan–Jun", "accounts", [0, 3, 8], NUM0, "Forecast months only; 2026 H1 is actual.", True)
    arow(7, "new_h2", "New paying investor firms per month, Jul–Dec", "accounts", [2, 5, 8], NUM0, "Used from Oct 2026 (Jul–Sep 2026 are actual).", True)
    arow(8, "churn", "Monthly logo churn", "% / month", [0.015, 0.015, 0.015], PCT, "Illustrative; pilot cohort too young to measure.", True)
    arow(9, "arpa_g", "ARPA growth (expansion)", "% / month", [0.01, 0.01, 0.01], PCT, "Seats and GateQ add-ons; NRR 112% reported Q3 2026.", True)
    section(11, "Cost of revenue and variable costs")
    arow(12, "cor", "Cost of revenue (hosting, model inference, payment fees)", "% of revenue", [0.22, 0.22, 0.22], PCT, "Gives the 78% gross margin reported for Q3 2026.", True)
    arow(13, "var", "Variable pilot onboarding and support", "% of revenue", [0.17, 0.17, 0.17], PCT, "Gives the 61% contribution margin reported for Q3 2026.", True)
    section(15, "People (monthly cost per head)")
    arow(16, "sal_f", "Founders", "USD / head / month", [1200, 2500, 3000], USD, "Founders take modest pay until the SAFE closes.")
    arow(17, "sal_e", "Engineering", "USD / head / month", [2000, 3000, 3300], USD, "Lagos market rates, illustrative.")
    arow(18, "sal_o", "Go-to-market and operations", "USD / head / month", [1500, 2200, 2400], USD, "Investor network, partnerships, customer success.")
    arow(19, "oncost", "Employer on-costs (pension, HMO, levies)", "% of salary", [0.10, 0.10, 0.10], PCT, "Illustrative.")
    section(21, "Fixed operating costs (per month)")
    arow(22, "fx_sw", "Software and tools", "USD / month", [600, 1500, 2500], USD, "")
    arow(23, "fx_office", "Office and co-working", "USD / month", [400, 1200, 1800], USD, "")
    arow(24, "fx_legal", "Legal, accounting and compliance", "USD / month", [300, 1500, 1500], USD, "")
    arow(25, "fx_travel", "Travel and investor-network events (West and East Africa)", "USD / month", [500, 4000, 6000], USD, "Use of funds: grow the investor network.")
    arow(26, "fx_mkt", "Marketing", "USD / month", [300, 3000, 5000], USD, "")
    arow(27, "fx_verif", "Verification and data-room partnerships", "USD / month", [0, 2500, 3000], USD, "Use of funds: Capital Q's own verification and data-room partnerships.")
    section(29, "Funding and cash")
    rows["open_cash"] = 30
    a["A30"], a["B30"], a["C30"] = "Opening cash, 1 Jan 2026 (founder funding)", "USD", 220000
    rows["safe_amt"] = 31
    a["A31"], a["B31"], a["C31"] = "Pre-seed SAFE amount", "USD", 1500000
    rows["safe_year"] = 32
    a["A32"], a["B32"], a["C32"] = "SAFE cash received — year", "year", 2027
    rows["safe_month"] = 33
    a["A33"], a["B33"], a["C33"] = "SAFE cash received — month", "month (1–12)", 3
    a["F31"] = "Matches companies.json raise target (USD 1,500,000, SAFE, pre-seed)."
    a["F33"] = "Target close 2027-03-31; modelled as received in March 2027."
    for r in range(30, 34):
        a.cell(row=r, column=1).font = F_N
        a.cell(row=r, column=2).font = F_N
        a.cell(row=r, column=3).font = F_IN
        a.cell(row=r, column=6).font = F_N
    a["C30"].number_format = USD
    a["C31"].number_format = USD
    a["C31"].fill = FILL_KEY
    section(35, "Actuals (self-reported demonstration data, Jan–Sep 2026)")
    a["A36"] = "Entered directly as blue inputs on 'Monthly Model' rows 'Paying investor firms' and 'MRR' for Jan–Sep 2026."
    a["A36"].font = F_N

    def A(key):  # absolute row range across the three year columns
        r = rows[key]
        return f"Assumptions!$C${r}:$E${r}"

    # ---- Monthly Model
    m = wb.create_sheet("Monthly Model")
    sheet_frame(m, "Monthly model, USD (Jan 2026 – Dec 2028)", 38)
    m.column_dimensions["A"].width = 46
    m.column_dimensions["B"].width = 12
    first_col = 3
    cols = [get_column_letter(first_col + i) for i in range(36)]
    for c in cols:
        m.column_dimensions[c].width = 10.5
    m.freeze_panes = "C6"
    head(m, 4, ["Line", "Unit"] + [f"{MON[mm - 1]} {str(y)[2:]}" for y, mm in MONTHS])
    R = {}
    labels = [
        ("year", "Year", ""),
        ("month", "Month", ""),
        ("type", "Actual / forecast", ""),
        (None, "", ""),
        ("hdr_rev", "REVENUE", ""),
        ("new", "New paying investor firms", "accounts"),
        ("acc", "Paying investor firms (end of month)", "accounts"),
        ("arpa", "ARPA (average revenue per account)", "USD"),
        ("mrr", "MRR (monthly recurring revenue)", "USD"),
        ("arr", "ARR run-rate (MRR × 12)", "USD"),
        (None, "", ""),
        ("hdr_m", "MARGINS", ""),
        ("cor", "Cost of revenue", "USD"),
        ("gp", "Gross profit", "USD"),
        ("gm", "Gross margin", "%"),
        ("var", "Variable onboarding and support", "USD"),
        ("cp", "Contribution", "USD"),
        ("cm", "Contribution margin", "%"),
        (None, "", ""),
        ("hdr_hc", "HEADCOUNT (end of month)", ""),
        ("hc_f", "Founders", "people"),
        ("hc_e", "Engineering", "people"),
        ("hc_o", "Go-to-market and operations", "people"),
        ("hc", "Total headcount", "people"),
        (None, "", ""),
        ("hdr_opex", "OPERATING COSTS", ""),
        ("pay", "Payroll incl. on-costs", "USD"),
        ("fx_sw", "Software and tools", "USD"),
        ("fx_office", "Office and co-working", "USD"),
        ("fx_legal", "Legal, accounting and compliance", "USD"),
        ("fx_travel", "Travel and investor-network events", "USD"),
        ("fx_mkt", "Marketing", "USD"),
        ("fx_verif", "Verification and data-room partnerships", "USD"),
        ("opex", "Total operating costs", "USD"),
        ("costs", "Total costs (cost of revenue + variable + operating)", "USD"),
        ("ebitda", "EBITDA", "USD"),
        ("burn", "Net burn (positive = cash consumed)", "USD"),
        (None, "", ""),
        ("hdr_cash", "CASH", ""),
        ("open", "Opening cash", "USD"),
        ("equity", "Equity / SAFE received", "USD"),
        ("close", "Closing cash", "USD"),
        ("runway", "Runway at 3-month average burn", "months"),
    ]
    r = 5
    for key, label, unit in labels:
        if key:
            R[key] = r
            c = m.cell(row=r, column=1, value=label)
            c.font = F_B if key.startswith("hdr") or key in ("mrr", "close", "burn") else F_N
            m.cell(row=r, column=2, value=unit).font = F_N
        r += 1

    def ref(key, col):
        return f"{col}{R[key]}"

    yi = lambda col: f"({col}${R['year']}-2025)"  # noqa: E731  (1..3)
    for i, (y, mm) in enumerate(MONTHS):
        col = cols[i]
        prev = cols[i - 1] if i else None
        actual = i < ACTUAL_MONTHS

        def put(key, value, fmt=USD, font=F_N):
            c = m.cell(row=R[key], column=first_col + i, value=value)
            c.number_format = fmt
            c.font = font
            return c

        put("year", y, "0", F_IN)
        put("month", mm, "0", F_IN)
        put("type", "Actual" if actual else "Forecast", "@", F_IN)
        if actual:
            put("new", None)
            put("acc", ACT_ACCOUNTS[i], NUM0, F_IN)
            put("mrr", ACT_MRR[i], USD, F_IN)
            put("arpa", f"=IF({ref('acc', col)}=0,0,{ref('mrr', col)}/{ref('acc', col)})", USD)
        else:
            put(
                "new",
                f"=IF({col}${R['month']}<=6,INDEX({A('new_h1')},{yi(col)}),INDEX({A('new_h2')},{yi(col)}))",
                NUM0,
                F_LINK,
            )
            put("acc", f"={prev}{R['acc']}*(1-INDEX({A('churn')},{yi(col)}))+{ref('new', col)}", NUM1)
            put("arpa", f"={prev}{R['arpa']}*(1+INDEX({A('arpa_g')},{yi(col)}))", USD)
            put("mrr", f"={ref('acc', col)}*{ref('arpa', col)}", USD)
        put("arr", f"={ref('mrr', col)}*12")
        put("cor", f"={ref('mrr', col)}*INDEX({A('cor')},{yi(col)})")
        put("gp", f"={ref('mrr', col)}-{ref('cor', col)}")
        put("gm", f"=IF({ref('mrr', col)}=0,0,{ref('gp', col)}/{ref('mrr', col)})", PCT)
        put("var", f"={ref('mrr', col)}*INDEX({A('var')},{yi(col)})")
        put("cp", f"={ref('gp', col)}-{ref('var', col)}")
        put("cm", f"=IF({ref('mrr', col)}=0,0,{ref('cp', col)}/{ref('mrr', col)})", PCT)
        f, e, o = headcount(i)
        put("hc_f", f, NUM0, F_IN)
        put("hc_e", e, NUM0, F_IN)
        put("hc_o", o, NUM0, F_IN)
        put("hc", f"=SUM({col}{R['hc_f']}:{col}{R['hc_o']})", NUM0)
        put(
            "pay",
            f"=({ref('hc_f', col)}*INDEX({A('sal_f')},{yi(col)})+{ref('hc_e', col)}*INDEX({A('sal_e')},{yi(col)})"
            f"+{ref('hc_o', col)}*INDEX({A('sal_o')},{yi(col)}))*(1+INDEX({A('oncost')},{yi(col)}))",
        )
        for k in ("fx_sw", "fx_office", "fx_legal", "fx_travel", "fx_mkt", "fx_verif"):
            put(k, f"=INDEX({A(k)},{yi(col)})", USD, F_LINK)
        put("opex", f"=SUM({col}{R['pay']}:{col}{R['fx_verif']})")
        put("costs", f"={ref('cor', col)}+{ref('var', col)}+{ref('opex', col)}")
        put("ebitda", f"={ref('mrr', col)}-{ref('costs', col)}")
        put("burn", f"=-{ref('ebitda', col)}")
        put("open", "=Assumptions!$C$30" if i == 0 else f"={prev}{R['close']}", USD, F_LINK if i == 0 else F_N)
        put(
            "equity",
            f"=IF(AND({col}${R['year']}=Assumptions!$C$32,{col}${R['month']}=Assumptions!$C$33),Assumptions!$C$31,0)",
        )
        put("close", f"={ref('open', col)}+{ref('equity', col)}-{ref('burn', col)}")
        lo = cols[max(0, i - 2)]
        put(
            "runway",
            f"=IF(AVERAGE({lo}{R['burn']}:{col}{R['burn']})<=0,\"n/a\",{ref('close', col)}/AVERAGE({lo}{R['burn']}:{col}{R['burn']}))",
            NUM1,
        )
    for key in ("mrr", "close"):
        for i in range(36):
            m.cell(row=R[key], column=first_col + i).font = F_B if i >= ACTUAL_MONTHS else Font(name=FONT, size=10, bold=True, color=BLUE)
    last_col = cols[-1]
    # ---- Annual Summary
    s = wb.create_sheet("Annual Summary")
    sheet_frame(s, "Annual summary and reconciliation to the Capital Q profile", 6)
    for col, w in zip("ABCDEF", (52, 16, 16, 16, 18, 44)):
        s.column_dimensions[col].width = w
    head(s, 4, ["USD unless stated", "FY2026", "FY2027", "FY2028"])
    yr = f"'Monthly Model'!$C${R['year']}:${last_col}${R['year']}"

    def mm_row(key):
        return f"'Monthly Model'!$C${R[key]}:${last_col}${R[key]}"

    srows = [
        ("Revenue", "sum", "mrr", USD),
        ("Cost of revenue", "sum", "cor", USD),
        ("Gross profit", "sum", "gp", USD),
        ("Gross margin", "ratio", ("gp", "mrr"), PCT),
        ("Contribution", "sum", "cp", USD),
        ("Contribution margin", "ratio", ("cp", "mrr"), PCT),
        ("Payroll incl. on-costs", "sum", "pay", USD),
        ("Total operating costs", "sum", "opex", USD),
        ("EBITDA", "sum", "ebitda", USD),
        ("Net burn", "sum", "burn", USD),
        ("Equity / SAFE received", "sum", "equity", USD),
        ("Closing cash (December)", "dec", "close", USD),
        ("Exit MRR (December)", "dec", "mrr", USD),
        ("Exit ARR run-rate (December)", "dec", "arr", USD),
        ("Paying investor firms (December)", "dec", "acc", NUM0),
        ("Headcount (December)", "dec", "hc", NUM0),
        ("Runway at December (months)", "dec", "runway", NUM1),
    ]
    for j, (label, kind, key, fmt) in enumerate(srows):
        row = 5 + j
        s.cell(row=row, column=1, value=label).font = F_N
        for k, y in enumerate(YEARS):
            col = get_column_letter(2 + k)
            if kind == "sum":
                v = f"=SUMIF({yr},{y},{mm_row(key)})"
            elif kind == "ratio":
                v = f"=IF(SUMIF({yr},{y},{mm_row(key[1])})=0,0,SUMIF({yr},{y},{mm_row(key[0])})/SUMIF({yr},{y},{mm_row(key[1])}))"
            else:
                v = f"=INDEX({mm_row(key)},{12 * (k + 1)})"
            c = s.cell(row=row, column=2 + k, value=v)
            c.number_format = fmt
            c.font = F_LINK
    rr = 5 + len(srows) + 2
    s.cell(row=rr, column=1, value="Reconciliation to the Capital Q profile (companies.json, as of 2026-09-30)").font = F_B
    head(s, rr + 1, ["Profile figure", "Profile value", "Model value", "", "Check", "Model source"])
    sep, jun = cols[8], cols[5]
    jul, aug = cols[6], cols[7]
    rec = [
        ("Monthly recurring revenue, Sep 2026 (USD)", 4800, f"='Monthly Model'!{sep}{R['mrr']}", USD, "Monthly Model, Sep 26"),
        ("ARR run-rate, end Sep 2026 (USD)", 57600, f"='Monthly Model'!{sep}{R['arr']}", USD, "MRR × 12"),
        ("Paying investor firms, Sep 2026", 12, f"='Monthly Model'!{sep}{R['acc']}", NUM0, "Monthly Model, Sep 26"),
        ("MRR growth Jun → Sep 2026", 0.60, f"='Monthly Model'!{sep}{R['mrr']}/'Monthly Model'!{jun}{R['mrr']}-1", PCT, "USD 3,000 → USD 4,800"),
        ("Gross margin, Q3 2026", 0.78, f"=SUM('Monthly Model'!{jul}{R['gp']}:{sep}{R['gp']})/SUM('Monthly Model'!{jul}{R['mrr']}:{sep}{R['mrr']})", PCT, "Jul–Sep 2026"),
        ("Contribution margin, Q3 2026", 0.61, f"=SUM('Monthly Model'!{jul}{R['cp']}:{sep}{R['cp']})/SUM('Monthly Model'!{jul}{R['mrr']}:{sep}{R['mrr']})", PCT, "Jul–Sep 2026"),
        ("Team size, Sep 2026", 6, f"='Monthly Model'!{sep}{R['hc']}", NUM0, "Headcount, Sep 26"),
        ("Pre-seed SAFE raise (USD)", 1500000, "=Assumptions!C31", USD, "Assumptions"),
    ]
    for j, (label, pv, mv, fmt, src) in enumerate(rec):
        row = rr + 2 + j
        s.cell(row=row, column=1, value=label).font = F_N
        c = s.cell(row=row, column=2, value=pv)
        c.font = F_IN
        c.number_format = fmt
        c = s.cell(row=row, column=3, value=mv)
        c.font = F_LINK
        c.number_format = fmt
        s.cell(row=row, column=5, value=f'=IF(ABS(C{row}-B{row})<0.0005*MAX(1,ABS(B{row})),"OK","CHECK")').font = F_B
        s.cell(row=row, column=6, value=src).font = F_N
    row = rr + 2 + len(rec) + 1
    notes = [
        "Reported but not modelled: net revenue retention 112% (Q3 2026); 1,240 active founders and investors (Sep 2026);",
        "estimated market USD 1.2 billion annual spend on deal sourcing and diligence tools (ESTIMATE). All self-reported, unverified.",
    ]
    for t in notes:
        s.cell(row=row, column=1, value=t).font = F_N
        row += 1
    wb.move_sheet("Annual Summary", offset=-1)
    wb.save(out)


# ---------------------------------------------------------------- cap table
def build_cap_table(out):
    wb = Workbook()
    props(wb, "Capital Q cap table", "Cap table, pre-round and pro forma for the proposed pre-seed SAFE")
    ws = wb.active
    ws.title = "Cap Table"
    sheet_frame(ws, "Capital Q — Cap table (illustrative), as of 30 September 2026", 8)
    for col, w in zip("ABCDEFGH", (40, 24, 16, 14, 18, 16, 14, 40)):
        ws.column_dimensions[col].width = w
    ws["A3"] = "Founder split, ESOP size and SAFE valuation cap are ILLUSTRATIVE placeholders for the demo; the owner confirms real terms."
    ws["A3"].font = Font(name=FONT, size=9, italic=True)
    head(ws, 5, ["Holder", "Role / instrument", "Shares (pre-round)", "% pre-round", "Shares (pro forma)", "% pro forma", "", "Note"])
    holders = [
        ("David Nwankwo", "Founder & CEO · ordinary shares", 6000000, "Illustrative split."),
        ("Captain Zino Mario", "Co-founder · ordinary shares", 3000000, "Illustrative split."),
        ("ESOP pool (unallocated)", "Reserved for employee options", 1000000, "Engineering hires draw from this pool."),
    ]
    for j, (h, role, sh, note) in enumerate(holders):
        r = 6 + j
        ws.cell(row=r, column=1, value=h).font = F_N
        ws.cell(row=r, column=2, value=role).font = F_N
        c = ws.cell(row=r, column=3, value=sh)
        c.font = F_IN
        c.number_format = NUM0
        c = ws.cell(row=r, column=4, value=f"=C{r}/$C$10")
        c.number_format = PCT
        c = ws.cell(row=r, column=5, value=f"=C{r}")
        c.number_format = NUM0
        c = ws.cell(row=r, column=6, value=f"=E{r}/$E$10")
        c.number_format = PCT
        ws.cell(row=r, column=8, value=note).font = F_N
    r = 9
    ws.cell(row=r, column=1, value="Pre-seed SAFE investors (proposed)").font = F_N
    ws.cell(row=r, column=2, value="Post-money SAFE, converts at next priced round").font = F_N
    ws.cell(row=r, column=3, value=0).number_format = NUM0
    ws.cell(row=r, column=4, value=f"=C{r}/$C$10").number_format = PCT
    c = ws.cell(row=r, column=5, value="=ROUND(SUM(C6:C8)*B17/(1-B17),0)")
    c.number_format = NUM0
    ws.cell(row=r, column=6, value=f"=E{r}/$E$10").number_format = PCT
    ws.cell(row=r, column=8, value="Shares shown as-if-converted at the cap.").font = F_N
    r = 10
    ws.cell(row=r, column=1, value="Total fully diluted").font = F_B
    for col, fmt in (("C", NUM0), ("D", PCT), ("E", NUM0), ("F", PCT)):
        c = ws[f"{col}10"]
        c.value = f"=SUM({col}6:{col}9)"
        c.number_format = fmt
        c.font = F_B
        c.border = Border(top=THIN)
    ws["H10"] = '=IF(AND(ROUND(D10,6)=1,ROUND(F10,6)=1),"Sums to 100%: OK","CHECK")'
    ws["H10"].font = F_B

    ws["A13"] = "Proposed round"
    ws["A13"].font = F_B
    terms = [
        ("Instrument", "SAFE (post-money valuation cap)", None, "Matches companies.json: instrument safe, stage pre-seed."),
        ("Amount (USD)", 1500000, USD, "Matches companies.json raise target."),
        ("Post-money valuation cap (USD)", 10000000, USD, "ILLUSTRATIVE — not in the Capital Q profile."),
        ("SAFE ownership at the cap", "=B15/B16", PCT, "Amount ÷ post-money cap."),
        ("Target close", "2027-03-31", None, "Matches companies.json target close date."),
    ]
    # Rows 14..18; B17 (ownership at the cap) drives the SAFE row above.
    for j, (label, v, fmt, note) in enumerate(terms):
        r = 14 + j
        ws.cell(row=r, column=1, value=label).font = F_N
        c = ws.cell(row=r, column=2, value=v)
        c.font = F_IN if not (isinstance(v, str) and v.startswith("=")) else F_N
        if fmt:
            c.number_format = fmt
        ws.cell(row=r, column=8, value=note).font = F_N
    ws["B16"].fill = FILL_KEY
    ws["A21"] = "No convertible notes, warrants or other instruments are outstanding in this illustration. Nothing here is a statutory register."
    ws["A21"].font = Font(name=FONT, size=9, italic=True)
    wb.save(out)


if __name__ == "__main__":
    out_dir = Path(sys.argv[1])
    out_dir.mkdir(parents=True, exist_ok=True)
    build_model(out_dir / "financial-model.xlsx")
    build_cap_table(out_dir / "capital-q-cap-table.xlsx")
